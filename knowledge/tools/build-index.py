#!/usr/bin/env python3
"""把 knowledge/知识库模板.xlsx 转成 entries/*.yaml 与 index/rules.json。

这是 knowledge/README.md 第 9 节那条流水线的实现：

    知识库模板.xlsx
       ↓ ① 解析 + 校验（本脚本）
    entries/*.yaml                ← 人读的、机器读的规则源文件
       ↓ ② 算向量 + 建索引（本脚本）
    index/rules.json              ← 浏览器运行时加载

设计约束（与 knowledge/schema.md 对齐）：

1. **单一真源是 xlsx。** entries/ 与 index/ 都是产物，脚本每次全量重写，
   不存在增量更新 —— 增量更新是「产物与真源不一致」的主要来源。
2. **校验失败一律中止**，报「sheet 名 + 行号 + 字段名 + 原因」。不写半份产物。
3. **错误与警告分开。** 错误中止；警告只列出来（schema.md 第 3 节）。
4. **向量在构建期算好**，运行时不批量推理。

依赖：openpyxl（读 xlsx）、PyYAML（写 yaml）。

用法：
    python knowledge/tools/build-index.py                 # 全量构建
    python knowledge/tools/build-index.py --check         # 只校验，不写文件
    python knowledge/tools/build-index.py --self-recall   # 构建后跑自召回冒烟测试
    python knowledge/tools/build-index.py --embedder bge  # 用真模型算向量
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import os
import re
import sys
import unicodedata
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path

# ---------------------------------------------------------------------------
# 依赖导入：缺依赖时给出「怎么装」，而不是抛出 ImportError 的堆栈
# ---------------------------------------------------------------------------

try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover - 环境问题，不参与单测
    sys.exit(
        "缺少 openpyxl，无法读取 xlsx。安装：\n"
        "    python -m pip install openpyxl PyYAML"
    )

# PyYAML 按需导入。只有真正读写 entries/*.yaml 的路径才需要它 —— `--check` 一个文件都不写，
# 不该因为缺一个 YAML 写入器就跑不起来。否则「这张表能不能过校验」会依赖一个与校验无关的库：
# 换台机器得先 pip install 才能自检，而自检恰恰是最该随时能跑的那一步。
_yaml_module = None


def require_yaml():
    """取 PyYAML 模块；首次调用时导入，并注册 LiteralStr 的块标量写法。"""
    global _yaml_module
    if _yaml_module is None:
        try:
            import yaml as _yaml
        except ImportError:  # pragma: no cover - 环境问题，不参与单测
            sys.exit(
                "缺少 PyYAML，无法读写 entries/*.yaml。安装：\n"
                "    python -m pip install openpyxl PyYAML"
            )
        _yaml.add_representer(LiteralStr, _represent_literal, Dumper=_yaml.SafeDumper)
        _yaml_module = _yaml
    return _yaml_module


# ---------------------------------------------------------------------------
# 路径与常量
# ---------------------------------------------------------------------------

ROOT = Path(__file__).resolve().parent.parent          # knowledge/
TEMPLATE = ROOT / "知识库模板.xlsx"
ENTRIES_DIR = ROOT / "entries"
INDEX_DIR = ROOT / "index"
REPORT_DIR = ROOT / "reports"

# 列顺序必须与 make-template.py 的 COLUMNS 一致。改动它等于改了团队的填写方式，
# 所以要同步改 knowledge/README.md 第 5 节与 schema.md 第 2 节。
COLUMNS = [
    "id",
    "标题",
    "适用学科",
    "适用字段",
    "触发条件",
    "规则内容",
    "正例",
    "反例",
    "来源",
    "权威度",
    "核验状态",
]

# 三个填写用 sheet。条目所属「类别」由它落在哪个 sheet 决定，
# 因此没有独立的类别列（schema.md 第 2 节的第一个设计决定）。
DATA_SHEETS = ("数据标准", "口径规则", "清洗案例")

# 模板里另外几张 sheet 是文档性内容，不是条目，明确跳过。
# 用显式白名单而不是「非 DATA_SHEETS 即跳过」：后者会让团队误建的 sheet 被静默忽略。
IGNORED_SHEETS = (
    "填写说明",
    "字段说明",
    # 2026-10-25 多学科扩展版带来的三张说明性 sheet：只给人读，不是条目。
    # 同样逐项登记而不是放宽成「未知 sheet 一律跳过」—— 拼错一个填写表的名字，
    # 应该当场报错，而不是悄悄少构建一批条目。
    "修订核验记录",
    "本次扩展导航",
    "学科目录",
)

# 学科 → 文件名 slug。条目落在 entries/<slug>.yaml。
# 不在表内的学科名会让构建失败并提示登记 —— 用拼音自动转换会产生
# 「同一个学科两种拼法」这类只能靠人发现的错。
SUBJECT_SLUGS = {
    "通用": "general",
    "心理学": "psychology",
    "教育学": "education",
    "公共卫生": "public-health",
    "社会学": "sociology",
    "经济学": "economics",
    "数据科学与大数据技术": "data-science",
}

AUTHORITY_LEVELS = ("国标", "行业", "团队经验")
VERIFY_STATES = ("已核对", "待核对")

# --- 权威度与来源必须匹配 ---
#
# 判据：依据能否追到一个成体系的方法学或官方文档规范。
# 命中 NATIONAL_STANDARD_SOURCES → 够得上「国标」（要求能给可核对的编号）；
# 命中 INDUSTRY_SOURCES → 「行业」（依据来自成体系的领域方法学，可引用而不必自证）；
# 都不命中 → 「团队经验」，含义是**无外部权威可引、由本团队采纳**，
# 而不是「我们跑过这批数据」—— 后者是「来源」列要讲的事，两者不要混。
#
# 新条目要标「行业」时，得先把它的依据类型加进下面的清单。
# 这是**刻意的摩擦**：一开始整批条目都默认被标成「行业」，等于这个档位没有信息量。
# 清单与 knowledge/schema.md 第 2 节那张表保持同步，改一处要改另一处。
NATIONAL_STANDARD_SOURCES = frozenset(
    {
        "行政区划代码国家标准（GB/T 2260；具体码值待团队核对原文）",
        "学历与学科专业代码国家标准（GB/T 4658 等；具体版本与码值待团队核对原文）",
    }
)

INDUSTRY_SOURCES = frozenset(
    {
        # 统计学与统计口径
        "统计学基本口径",
        "统计分析通用口径",
        "统计推断的通用口径",
        "统计分组的通用口径",
        "统计报表的通用口径",
        "统计指标口径的通用约定",
        "统计数据清洗的通用口径",
        "数据呈现的通用口径",
        "计数数据处理的通用口径",
        "缺失数据处理的通用口径",
        "抽样调查的通用口径",
        "面板数据处理的通用口径",
        # 领域方法学
        "心理测量学通用口径",
        "测量学与统计的通用口径",
        "教育统计的通用口径",
        "教育统计的通用口径（具体代码待团队对照现行标准原文）",
        "国内社会科学调查数据的常见约定；具体码表待按数据集说明文档核对",
        "身份信息字段的通用校验约定（具体校验算法待团队对照权威口径）",
        "编码标准存在版本差异这一已知事实（具体标准号与码值待团队对照原文）",
        # 国标体系、但具体编号尚未核对 —— 核对原文后应升为「国标」
        "民族编码相关国家标准（具体标准号与码值待团队核对原文）",
        # 工程方法学
        "数值计算的通用约定",
        "大数据工程实践的常见问题",
        "半结构化数据处理的常见问题",
        "数据管道的常见问题",
        "数据处理留痕的通用要求",
    }
)

ID_PATTERN = re.compile(r"^[a-z0-9-]+$")

# 「来源」里出现一个可点开的 URL，本身就是可引用证据。
#
# 2026-10-25 团队交付的多学科扩展版把来源写成「依据名称 + URL + 核对日期」，
# 每条都能追到公开文档（pandas / Microsoft / NIST / Unicode / RFC / 官方标准平台…）。
# 继续要求「来源逐字命中 INDUSTRY_SOURCES」会把这 120 条全报成警告，
# 而警告一旦变成噪音，真正的漂移就没人看了。
# 所以判据改成：命中旧清单 **或** 来源里带 URL —— 后者可核对，才够得上「行业」。
# 旧清单保留：它让不联网、无 URL 的历史条目照旧可校验。
SOURCE_URL_PATTERN = re.compile(r"https?://\S+")

# 国标的判据原本是「来源逐字命中清单」；新库把标准写进了来源文本
#（如「国家标准信息公共服务平台：GB/T 4658-2006…」）。要守住的是**可核对的编号**，
# 所以改成「来源里出现 GB 编号」也算数 —— 编号在，就能去平台查。
# 清单仍保留，兼容旧写法。
GB_NUMBER_PATTERN = re.compile(r"GB(?:/T)?\s*\d{3,5}(?:\s*[—\-–]\s*\d{2,4})?")


# 「规则内容含同上/略/见上」这条校验要窄化：裸匹配「略」会命中「策略」「忽略」
# 「省略」这些正常用词。真正要拦的是「此处略」这类占位写法。
#
# 收窄踩过两次坑，轨迹记下来：
#   1. 裸匹配「略」→ 命中「策略」等正常词
#   2. 只匹配「略。」→ 仍然命中「再决定填补策略。」（实际被构建拦下的就是我们自己写的一条规则）
#   3. 要求「略」前面不是汉字 → 「此处略。」这种真占位反而漏掉了
# 现在的做法：只排除「以略结尾的常见词」的最后一个字。这个清单要跟着语言习惯维护，
# 但比前两种都准 —— 加词的代价远小于漏报或误报。
LUECED_COMPOUND_TAILS = "策战省忽简大约概领侵谋方韬经节"
PLACEHOLDER_PATTERN = re.compile(
    rf"同上|见上|（略）|\(略\)|(?<![{LUECED_COMPOUND_TAILS}])略。|^\s*略\s*$"
)

# 「归并类规则必须有反例」。识别方式是在标题与规则内容里找归并语义。
# 宁可多识别（多要求一条反例）也不漏 —— 缺反例是误合并的直接来源。
MERGE_HINT_PATTERN = re.compile(
    r"归并|归为|同义|合并|等价|写法不统一|统一为|统一成|应视为同一|视为同一取值"
)

# dense 向量维度。schema.md 第 4 节的示例是 512 维（bge-small-zh-v1.5 的维度），
# 这里保持 512，让两种后端产出同形状的产物，运行时无需分支。
EMBEDDING_DIM = 512

# 向量后端的标识符。放在这里是因为下面的阈值表要按后端取值 ——
# 常量区不能在文件后面才定义，否则定义顺序就成了隐藏依赖。
EMBEDDER_HASH = "char-ngram-hash"
EMBEDDER_BGE = "bge"

# 检索参数。
RRF_K = 60            # 倒数排名融合的平滑常数，RRF 原论文取值
BM25_K1 = 1.2         # 词频饱和参数，Robertson 标准取值
BM25_B = 0.75         # 长度归一化强度，同上
FIELD_HIT_BOOST = 0.05  # 列名命中 适用字段 的加权
TOP_K = 5             # 注入 Prompt 的条目数上限

# 近重复判定的余弦阈值。**必须按后端分别给**：
# 语义模型的余弦天然偏高（同义句常 0.9+），而哈希 2-gram 的余弦尺度被压得很扁 ——
# 实测 96 条 4560 对里最大只有 0.38。用同一个阈值会让离线后端下这个指标永远报 0。
#
# 0.30 是照当前离线后端的分布定的（p99 = 0.19、最大 0.38），属于**未经标注集校准**的经验值。
# 换 bge 后端后必须重定，Plan 4 建起评测集后应重新扫。
NEAR_DUPLICATE_COSINE = {
    EMBEDDER_HASH: 0.30,
    EMBEDDER_BGE: 0.86,
}

# 报告里无条件展示的「最相似 TOP N」。即使一对都没过阈值也要展示 ——
# 否则读者无法区分「真的没有重复」与「阈值根本够不着」。
SIMILARITY_TOP_N = 5


# ---------------------------------------------------------------------------
# 文本规范化
#
# 与 src/core/normalize/index.ts 保持同一套语义（NFKC 覆盖全角→半角，
# 去零宽字符），这样「产品里的规范化」与「索引里的规范化」不会漂移。
# ---------------------------------------------------------------------------

INVISIBLE_CHARS = re.compile(
    "[\u200b-\u200f\u202a-\u202e\u2060-\u2064\ufeff]"
)


def normalize_text(text: str) -> str:
    """去零宽字符 → NFKC（全角转半角）→ 去首尾空白 → 压缩内部连续空白。"""
    cleaned = INVISIBLE_CHARS.sub("", text)
    cleaned = unicodedata.normalize("NFKC", cleaned)
    return re.sub(r"\s+", " ", cleaned).strip()


def to_grams(text: str) -> list[str]:
    """字符 2-gram。

    先去掉所有空白再切 gram，否则「硕士 研究生」会切出「士 」「 研」这种
    带空格的 gram，与「硕士研究生」召不回同一条规则。
    """
    compact = re.sub(r"\s+", "", normalize_text(text).lower())
    if len(compact) < 2:
        return [compact] if compact else []
    return [compact[i : i + 2] for i in range(len(compact) - 1)]


def term_frequency(text: str) -> dict[str, int]:
    counts: dict[str, int] = {}
    for gram in to_grams(text):
        counts[gram] = counts.get(gram, 0) + 1
    return counts


def split_aliases(raw: str) -> list[str]:
    """把「学历,文化程度，最高学历」拆成别名列表，去掉空白与重复（保序）。"""
    if not raw:
        return []
    parts = re.split(r"[,，、;；]", raw)
    seen: set[str] = set()
    out: list[str] = []
    for part in parts:
        alias = normalize_text(part)
        if not alias or alias in seen:
            continue
        seen.add(alias)
        out.append(alias)
    return out


# ---------------------------------------------------------------------------
# 解析层：xlsx → Entry 列表
# ---------------------------------------------------------------------------


@dataclass
class Diagnostic:
    """一条校验诊断。level 为 error 时构建中止。"""

    level: str          # error | warning
    sheet: str
    row: int | None     # Excel 里的真实行号（1 基，含表头）
    entry_id: str
    field: str
    message: str

    def format(self) -> str:
        where = f"{self.sheet} 第 {self.row} 行" if self.row else self.sheet
        target = f"[{self.entry_id}] " if self.entry_id else ""
        return f"{where} · {self.field} · {target}{self.message}"


@dataclass
class Entry:
    """一条规则。字段名与 Excel 列名一致，便于互相追溯。"""

    id: str
    标题: str = ""
    类别: str = ""
    适用学科: str = "通用"
    适用字段: list[str] = field(default_factory=list)
    触发条件: str = ""
    规则内容: str = ""
    正例: str = ""
    反例: str = ""
    来源: str = ""
    权威度: str = "行业"
    核验状态: str = "待核对"
    # 溯源信息（不进 yaml，只进报告，便于团队按行号回查 Excel）
    sheet: str = ""
    row: int = 0

    def to_yaml_dict(self) -> dict:
        """按 schema.md 第 4 节的键顺序输出。"""
        return {
            "id": self.id,
            "标题": self.标题,
            "类别": self.类别,
            "适用学科": self.适用学科,
            "适用字段": self.适用字段,
            "触发条件": self.触发条件,
            "规则内容": self.规则内容,
            "正例": self.正例,
            "反例": self.反例,
            "来源": self.来源,
            "权威度": self.权威度,
            "核验状态": self.核验状态,
        }

    def full_text(self) -> str:
        """条目全文，用于稀疏检索。"""
        return " ".join(
            [
                self.标题,
                " ".join(self.适用字段),
                self.触发条件,
                self.规则内容,
                self.正例,
                self.反例,
            ]
        )

    def dense_text(self) -> str:
        """稠密检索用的文本。

        按 schema.md 第 4 节的约定只拼「标题 + 触发条件 + 规则内容」——
        正例/反例里含大量具体取值，混进稠密文本会把语义中心带偏。
        """
        return "\n".join([self.标题, self.触发条件, self.规则内容])


def cell_to_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    if isinstance(value, float) and value.is_integer():
        # Excel 把纯数字文本存成 float 是常态（如 id 列写成 123）
        return str(int(value))
    return str(value).strip()


def load_entries(template: Path, diagnostics: list[Diagnostic]) -> list[Entry]:
    """读取三个填写 sheet。表头不符会直接报错 —— 那是模板被改动的信号。"""
    if not template.exists():
        sys.exit(
            f"找不到模板：{template}\n"
            "模板由 knowledge/tools/make-template.py 生成。"
        )

    workbook = load_workbook(template, read_only=True, data_only=True)
    sheet_names = list(workbook.sheetnames)

    unknown = [
        name
        for name in sheet_names
        if name not in DATA_SHEETS and name not in IGNORED_SHEETS
    ]
    for name in unknown:
        diagnostics.append(
            Diagnostic(
                "error",
                name,
                None,
                "",
                "sheet 名",
                f"不是三个已知类别之一，也未登记为文档性 sheet。"
                f"条目类别无法推断。已知类别：{'、'.join(DATA_SHEETS)}；"
                f"已知文档 sheet：{'、'.join(IGNORED_SHEETS)}。"
                f"新增填写表请同步改本脚本与 schema.md。",
            )
        )

    entries: list[Entry] = []

    for sheet_name in DATA_SHEETS:
        if sheet_name not in sheet_names:
            diagnostics.append(
                Diagnostic(
                    "error",
                    sheet_name,
                    None,
                    "",
                    "sheet",
                    "模板里缺少这张填写表。请重新运行 make-template.py 或检查改名。",
                )
            )
            continue

        sheet = workbook[sheet_name]
        rows = sheet.iter_rows(values_only=True)

        try:
            header = next(rows)
        except StopIteration:
            diagnostics.append(
                Diagnostic("error", sheet_name, 1, "", "表头", "空 sheet，没有表头行")
            )
            continue

        header_text = [cell_to_text(c) for c in header]
        if header_text[: len(COLUMNS)] != COLUMNS:
            diagnostics.append(
                Diagnostic(
                    "error",
                    sheet_name,
                    1,
                    "",
                    "表头",
                    f"列名与约定不一致。期望 {COLUMNS}，实际 {header_text[: len(COLUMNS)]}",
                )
            )
            continue

        for offset, row in enumerate(rows):
            row_number = offset + 2  # Excel 里没有 0 行，且第 1 行是表头
            # 模板预留给团队 40 个带边框的空行，必须跳过；否则会产出 40 条空条目
            if all(not cell_to_text(c) for c in row):
                continue

            values = [cell_to_text(c) for c in row[: len(COLUMNS)]]
            record = dict(zip(COLUMNS, values))

            entry = Entry(
                id=record["id"],
                标题=record["标题"],
                类别=sheet_name,
                适用学科=record["适用学科"] or "通用",
                适用字段=split_aliases(record["适用字段"]),
                触发条件=record["触发条件"],
                规则内容=record["规则内容"],
                正例=record["正例"],
                反例=record["反例"],
                来源=record["来源"],
                权威度=record["权威度"] or "行业",
                核验状态=record["核验状态"] or "待核对",
                sheet=sheet_name,
                row=row_number,
            )
            entries.append(entry)

    workbook.close()
    return entries


# ---------------------------------------------------------------------------
# 校验层：schema.md 第 3 节
# ---------------------------------------------------------------------------


def validate(entries: list[Entry], diagnostics: list[Diagnostic]) -> None:
    """错误一律中止；警告只记录。"""

    seen_ids: dict[str, Entry] = {}
    alias_owners: dict[str, list[str]] = {}
    # 判定「别名冲突」时要知道各条目属于哪个类别，用于区分「同类别内冗余」与「跨类别分工」
    category_of: dict[str, str] = {entry.id: entry.类别 for entry in entries if entry.id}

    for entry in entries:
        where = (entry.sheet, entry.row, entry.id)

        # --- 错误 ---
        if not entry.id:
            diagnostics.append(
                Diagnostic("error", *where[:2], "", "id", "为空")
            )
        elif not ID_PATTERN.match(entry.id):
            diagnostics.append(
                Diagnostic(
                    "error", *where[:2], entry.id, "id",
                    "含非法字符。只允许小写字母、数字与短横线（会出现在索引文件名与 URL 里）",
                )
            )
        elif entry.id in seen_ids:
            first = seen_ids[entry.id]
            diagnostics.append(
                Diagnostic(
                    "error", *where[:2], entry.id, "id",
                    f"重复（首次出现在 {first.sheet} 第 {first.row} 行）。"
                    f"后一条会静默覆盖前一条",
                )
            )
        else:
            seen_ids[entry.id] = entry

        for field_name in ("标题", "触发条件", "规则内容", "来源"):
            if not getattr(entry, field_name):
                diagnostics.append(
                    Diagnostic("error", *where[:2], entry.id, field_name, "为空")
                )

        if entry.来源 and entry.来源 in ("待确认", "待核对"):
            # 不是错误：「不确定」是允许的状态（schema.md 第 2 节）
            diagnostics.append(
                Diagnostic(
                    "warning", *where[:2], entry.id, "来源",
                    f"填的是「{entry.来源}」。允许，但答辩无法追溯到具体出处，"
                    f"建议确认后换成国标号／口径名称／「我们自己跑数据发现」",
                )
            )

        if entry.规则内容 and PLACEHOLDER_PATTERN.search(entry.规则内容):
            diagnostics.append(
                Diagnostic(
                    "error", *where[:2], entry.id, "规则内容",
                    "含「同上／略／见上」类占位写法。检索到的片段会脱离上下文而不自洽",
                )
            )

        if entry.权威度 not in AUTHORITY_LEVELS:
            diagnostics.append(
                Diagnostic(
                    "error", *where[:2], entry.id, "权威度",
                    f"「{entry.权威度}」不在枚举内，只能是 {'／'.join(AUTHORITY_LEVELS)}",
                )
            )

        if entry.核验状态 not in VERIFY_STATES:
            diagnostics.append(
                Diagnostic(
                    "error", *where[:2], entry.id, "核验状态",
                    f"「{entry.核验状态}」不在枚举内，只能是 {'／'.join(VERIFY_STATES)}",
                )
            )

        if MERGE_HINT_PATTERN.search(entry.标题 + entry.规则内容) and not entry.反例:
            diagnostics.append(
                Diagnostic(
                    "error", *where[:2], entry.id, "反例",
                    "归并类规则必须给反例。缺反例是误合并的主要来源",
                )
            )

        # 权威度与来源必须对得上。只报警告不中止 —— 它是给人复核的提示，
        # 不是机器能替人下的判断（「这条够不够得上行业」最终要人来定）。
        if entry.权威度 == "国标" and not (
            entry.来源 in NATIONAL_STANDARD_SOURCES or GB_NUMBER_PATTERN.search(entry.来源)
        ):
            diagnostics.append(
                Diagnostic(
                    "warning", *where[:2], entry.id, "权威度",
                    f"标了「国标」但来源「{entry.来源}」里既不在国标依据清单、也没有可核对的 GB 编号。"
                    f"请补上标准号，或改标「行业」",
                )
            )
        elif entry.权威度 == "行业" and not (
            entry.来源 in INDUSTRY_SOURCES or SOURCE_URL_PATTERN.search(entry.来源)
        ):
            diagnostics.append(
                Diagnostic(
                    "warning", *where[:2], entry.id, "权威度",
                    f"标了「行业」但来源「{entry.来源}」既不在行业依据清单、也没有可核对的 URL。"
                    f"够得上行业就补一个可点开的来源链接（或把依据类型加进 build-index.py 的 INDUSTRY_SOURCES），"
                    f"够不上就改标「团队经验」",
                )
            )
        elif (
            entry.权威度 == "团队经验"
            and entry.来源 in INDUSTRY_SOURCES
        ):
            diagnostics.append(
                Diagnostic(
                    "warning", *where[:2], entry.id, "权威度",
                    f"标了「团队经验」但来源「{entry.来源}」是可引用的方法学依据，够得上「行业」",
                )
            )

        # --- 篇幅 ---
        if len(entry.标题) > 30:
            diagnostics.append(
                Diagnostic(
                    "warning", *where[:2], entry.id, "标题",
                    f"{len(entry.标题)} 字，超过 30 字上限",
                )
            )
        if len(entry.触发条件) > 60:
            diagnostics.append(
                Diagnostic(
                    "warning", *where[:2], entry.id, "触发条件",
                    f"{len(entry.触发条件)} 字，超过 60 字上限",
                )
            )
        sentences = len([s for s in re.split(r"[。！？；;]", entry.规则内容) if s.strip()])
        if entry.规则内容 and not (2 <= sentences <= 5):
            diagnostics.append(
                Diagnostic(
                    "warning", *where[:2], entry.id, "规则内容",
                    f"{sentences} 句，约定 2–5 句。太少说不清边界，太多检索到也不自洽",
                )
            )

        if entry.适用学科 not in SUBJECT_SLUGS:
            diagnostics.append(
                Diagnostic(
                    "error", *where[:2], entry.id, "适用学科",
                    f"「{entry.适用学科}」未登记 slug，无法决定写进哪个 yaml。"
                    f"请在 build-index.py 的 SUBJECT_SLUGS 里登记（已登记："
                    f"{'、'.join(SUBJECT_SLUGS)}）",
                )
            )

        if len(entry.适用字段) == 1:
            diagnostics.append(
                Diagnostic(
                    "warning", *where[:2], entry.id, "适用字段",
                    f"只有「{entry.适用字段[0]}」一个别名。同一件事不同人的表头写法不同"
                    f"（学历／文化程度／最高学历／edu），别名写全才能命中",
                )
            )
        # 留空不报警告：那是 schema.md 第 2 节明确设计的「泛化规则」语义。
        # 给它报警告会让 8 条合法条目各出一条噪音，把真正的信号（下面的跨行别名冲突）淹掉。
        # 泛化规则的数量在报告与 field-alias-taxonomy.json 里单独统计。

        for alias in entry.适用字段:
            alias_owners.setdefault(alias, []).append(entry.id)

    # 同一别名被多条规则声明 → 可能重复规则。
    #
    # 但要分成两类看，否则警告会淹掉真信号（47 条规则时这里能刷出 32 条）：
    #   · 同类别内共用 —— 两条都属「数据标准」管着同一个字段，冗余的可能性高，值得逐条确认
    #   · 跨类别共用 —— 一条「数据标准」加一条「清洗案例」都覆盖该字段，这正是约定的分工
    #     （标准讲取值规范、案例讲踩过的坑），属正常，只汇总一行报告数量
    # 真正重复的规则由近重复检测（余弦）负责，这里只管「它们会争同一批列」这件事。
    same_category: list[tuple[str, list[str]]] = []
    cross_category: list[tuple[str, list[str]]] = []

    for alias, owners in sorted(alias_owners.items()):
        if len(owners) < 2:
            continue
        categories = {category_of.get(owner, "?") for owner in owners}
        if len(categories) == 1:
            same_category.append((alias, owners))
        else:
            cross_category.append((alias, owners))

    for alias, owners in same_category:
        labelled = "、".join(f"{owner}（{category_of.get(owner, '?')}）" for owner in owners)
        diagnostics.append(
            Diagnostic(
                "warning",
                "（跨行）",
                None,
                "",
                "适用字段",
                f"别名「{alias}」被同属「{category_of.get(owners[0], '?')}」的 "
                f"{len(owners)} 条规则声明：{labelled}。两条规则管同一个字段时冗余的可能性较高，"
                f"请确认是互补规则还是重复规则",
            )
        )

    if cross_category:
        diagnostics.append(
            Diagnostic(
                "warning",
                "（跨行）",
                None,
                "",
                "适用字段",
                f"另有 {len(cross_category)} 个别名被跨类别的多条规则共用"
                f"（{'、'.join(alias for alias, _ in cross_category)}）。"
                f"这是约定的分工 —— 标准讲取值规范、口径讲怎么算、案例讲踩过的坑，"
                f"同一字段由三处分别覆盖是对的。仅在其内容确实重复时才需合并",
            )
        )


# ---------------------------------------------------------------------------
# 向量层
#
# 两个后端：
#   char-ngram-hash  离线、确定性、无额外依赖。产出的是真实的稠密向量
#                    （哈希 2-gram 词袋 + L2 归一化），但不是语义模型。
#   bge              真模型 BAAI/bge-small-zh-v1.5，需要 sentence-transformers。
#
# 为什么要有一个离线后端：schema.md 第 4 节要求向量在构建期算好，
# 而构建不该依赖网络与 200MB 模型下载。缺模型时给出「一个能跑的降级路径」
# 而不是「构建失败」，这与项目其它部分（端侧 embedding 不可用 → 降级结构化规则）
# 的处理方式一致。model 字段如实写出用了哪个后端，不冒名。
# （后端标识符 EMBEDDER_HASH / EMBEDDER_BGE 定义在文件顶部的常量区。）
# ---------------------------------------------------------------------------


def embed_hash_ngram(text: str, dim: int = EMBEDDING_DIM) -> list[float]:
    """哈希字符 2-gram 词袋，L2 归一化。

    用 blake2b 而不是内置 hash()：CPython 对 str 的 hash 每个进程都不同
    （PYTHONHASHSEED 随机化），用它算向量会让「同一份输入两次构建产出不同索引」，
    可复现性直接失效。
    """
    vector = [0.0] * dim
    for gram, count in term_frequency(text).items():
        digest = hashlib.blake2b(gram.encode("utf-8"), digest_size=8).digest()
        bucket = int.from_bytes(digest[:4], "big") % dim
        # 第 5 个字节决定符号，抵消哈希碰撞带来的系统性正偏
        sign = 1.0 if digest[4] & 1 else -1.0
        vector[bucket] += sign * (1.0 + math.log(count))

    norm = math.sqrt(sum(value * value for value in vector))
    if norm == 0.0:
        return vector
    return [value / norm for value in vector]


class BgeEmbedder:
    """真语义后端。模型不可用时抛 RuntimeError，由调用方决定怎么处理。"""

    name = "BAAI/bge-small-zh-v1.5"

    def __init__(self) -> None:
        try:
            from sentence_transformers import SentenceTransformer
        except ImportError as exc:  # pragma: no cover - 取决于本机环境
            raise RuntimeError(
                "缺少 sentence-transformers。安装：\n"
                "    python -m pip install sentence-transformers\n"
                "或改用离线后端：--embedder char-ngram-hash"
            ) from exc

        self._model = SentenceTransformer(self.name)

    def encode(self, texts: list[str]) -> list[list[float]]:
        # BGE 系列检索时应给查询加指令前缀，文档侧不加
        vectors = self._model.encode(texts, normalize_embeddings=True)
        return [[float(value) for value in row] for row in vectors]


def embed_all(
    texts: list[str], backend: str
) -> tuple[list[list[float]], str]:
    """返回 (向量列表, 实际使用的 model 标识)。"""
    if backend == EMBEDDER_BGE:
        embedder = BgeEmbedder()
        return embedder.encode(texts), embedder.name

    return (
        [embed_hash_ngram(text) for text in texts],
        f"{EMBEDDER_HASH}-{EMBEDDING_DIM}-v1（离线降级后端，非语义模型；"
        f"换语义模型：--embedder bge）",
    )


def cosine(left: list[float], right: list[float]) -> float:
    dot = sum(a * b for a, b in zip(left, right))
    # 两个向量在产出时都已 L2 归一化，这里不再重复算模长
    return max(-1.0, min(1.0, dot))


# ---------------------------------------------------------------------------
# 输出层：entries/*.yaml
# ---------------------------------------------------------------------------


class LiteralStr(str):
    """标记为「用 | 块写出来」的字符串，纯粹为了 git diff 好读。"""


def _represent_literal(dumper, data: LiteralStr):
    """把 LiteralStr 写成 | 块。注册动作在 require_yaml() 里，理由见那里。"""
    return dumper.represent_scalar("tag:yaml.org,2002:str", str(data), style="|")


def _as_literal(text: str) -> str | LiteralStr:
    """长文本用 | 块；短的、含游离空格的用默认引号写法。

    `|` 块会原样保留换行，因此文本里若有「行尾空格」或首尾空白，
    往返一趟就可能变样 —— 这类一律不用块写法，避免静默改内容。
    """
    if not text or len(text) < 40:
        return text
    if text != text.strip():
        return text
    if any(line != line.rstrip() for line in text.split("\n")):
        return text
    return LiteralStr(text)


def dump_entries_yaml(entries: list[Entry], subject: str, slug: str) -> str:
    header = (
        f"# 本文件由 knowledge/tools/build-index.py 生成，请勿手改。\n"
        f"# 真源是 knowledge/知识库模板.xlsx；手改的内容会在下次构建时被覆盖。\n"
        f"# 学科：{subject}　条目数：{len(entries)}\n"
        f"#\n"
        f"# 字段定义见 knowledge/schema.md 第 4 节。\n"
        f"\n"
    )

    payload = []
    for entry in entries:
        row = entry.to_yaml_dict()
        payload.append(
            {
                "id": row["id"],
                "标题": _as_literal(row["标题"]),
                "类别": row["类别"],
                "适用学科": row["适用学科"],
                "适用字段": row["适用字段"],
                "触发条件": _as_literal(row["触发条件"]),
                "规则内容": _as_literal(row["规则内容"]),
                "正例": _as_literal(row["正例"]),
                "反例": _as_literal(row["反例"]),
                "来源": _as_literal(row["来源"]),
                "权威度": row["权威度"],
                "核验状态": row["核验状态"],
            }
        )

    body = require_yaml().safe_dump(
        payload,
        allow_unicode=True,
        sort_keys=False,
        default_flow_style=False,
        width=88,
        indent=2,
    )

    # 自证：块标量在某些内容下会悄悄改变取值。重载一遍比对，
    # 不一致就退回默认引号写法 —— 宁可不漂亮，也不能改内容。
    try:
        reloaded = require_yaml().safe_load(body)
        expected = [{**row, "标题": str(row["标题"])} for row in payload]
        actual = [{**item, "标题": str(item["标题"])} for item in reloaded]
        if actual != expected:
            raise ValueError("literal block round-trip mismatch")
    except Exception:  # noqa: BLE001 - 任何往返问题都退回安全写法
        body = require_yaml().safe_dump(
            payload,
            allow_unicode=True,
            sort_keys=False,
            default_flow_style=False,
            width=10000,
        )

    return header + body


# ---------------------------------------------------------------------------
# 输出层：index/rules.json
# ---------------------------------------------------------------------------


def build_rules_json(
    entries: list[Entry],
    embeddings: list[list[float]],
    model: str,
    built_at: str,
) -> dict:
    records = []
    for entry, vector in zip(entries, embeddings):
        records.append(
            {
                "id": entry.id,
                "学科": [entry.适用学科],
                "适用字段": entry.适用字段,
                # 条目原文。schema.md 的示例只画了 id/学科/适用字段/embedding/sparse，
                # 但同段文字写明产物是「条目原文 + 稠密向量 + 稀疏词频」——
                # 运行时若拿不到正文，检索到也没东西可注入 Prompt。
                "entry": entry.to_yaml_dict(),
                "embedding": [round(value, 6) for value in vector],
                "sparse": term_frequency(entry.full_text()),
            }
        )

    return {
        "builtAt": built_at,
        "model": model,
        "dim": EMBEDDING_DIM,
        "rrfK": RRF_K,
        "topK": TOP_K,
        "entries": records,
    }


def write_json(path: Path, payload: object, compact_number_arrays: bool = False) -> None:
    """写 UTF-8（无 BOM）、LF、2 空格缩进的 JSON。

    `compact_number_arrays` 把「只含数字的数组」压成一行。512 维向量若按
    indent=2 展开是 512 行，17 条就是 8704 行 —— 那样的产物没法读、
    git diff 也看不出任何有意义的变化。压成一行后，向量之间的差异
    仍然可见（数值变了那一行就变了），文件从 200 KB 降到 40 KB 量级。
    """
    text = json.dumps(payload, ensure_ascii=False, indent=2, sort_keys=False)
    if compact_number_arrays:
        text = _collapse_number_arrays(text)
    path.write_text(text + "\n", encoding="utf-8", newline="\n")


NUMBER_LINE = re.compile(r"^\s*-?\d+(?:\.\d+)?(?:[eE][-+]?\d+)?,?\s*$")


def _collapse_number_arrays(text: str) -> str:
    """把 JSON 文本里「元素全是数字」的多行数组合并成一行。

    刻意只碰数字行：含字符串的数组（如 `适用字段`）保持多行，
    因为那些列表要能一眼看清有哪几个别名。
    """
    lines = text.split("\n")
    out: list[str] = []
    index = 0
    while index < len(lines):
        current = lines[index]
        if not current.rstrip().endswith("["):
            out.append(current)
            index += 1
            continue

        numbers: list[str] = []
        cursor = index + 1
        closing = None
        while cursor < len(lines):
            candidate = lines[cursor]
            stripped = candidate.strip()
            if stripped in ("]", "],"):
                closing = stripped
                break
            if not NUMBER_LINE.match(candidate):
                # 不是纯数字数组（例如字符串列表），整段原样保留
                break
            numbers.append(stripped.rstrip(","))
            cursor += 1

        if closing is None or not numbers:
            out.append(current)
            index += 1
            continue

        out.append(current + ", ".join(numbers) + "]" + ("," if closing == "]," else ""))
        index = cursor + 1

    return "\n".join(out)


def verify_rules_json(path: Path, expected_count: int, dim: int) -> None:
    """写完再读回来验一遍，不合格就抛错中止。

    为什么非要读回来：产物是构建期一次性算出来的，坏了要等到运行时（或 Plan 3
    把检索搬进浏览器之后）才暴露。而这里有两处真会弄坏它 —— 规则正文进 JSON 字符串
    要转义引号／换行／反斜杠（`general-fullwidth` 的正文里就有字面量 `\\r`），
    向量数组又被 `_collapse_number_arrays` 重写过。yaml 那边有往返自证，
    json 这边也得有。
    """
    payload = json.loads(path.read_text(encoding="utf-8"))
    records = payload.get("entries")
    if not isinstance(records, list):
        raise ValueError("rules.json 缺少 entries 数组")
    if len(records) != expected_count:
        raise ValueError(
            f"rules.json 条目数不符：期望 {expected_count}，实际 {len(records)}"
        )
    if payload.get("dim") != dim:
        raise ValueError(f"rules.json 的 dim 应为 {dim}，实际 {payload.get('dim')}")

    for record in records:
        entry_id = record.get("id", "<无 id>")

        vector = record.get("embedding")
        if not isinstance(vector, list) or len(vector) != dim:
            raise ValueError(f"[{entry_id}] 向量维度不对：{len(vector) if isinstance(vector, list) else '不是数组'}")
        if not all(isinstance(value, (int, float)) for value in vector):
            raise ValueError(f"[{entry_id}] 向量里混进了非数字")
        norm = math.sqrt(sum(value * value for value in vector))
        if abs(norm - 1.0) > 1e-3:
            raise ValueError(f"[{entry_id}] 向量未归一化，模长 {norm:.6f}")

        if not record.get("sparse"):
            raise ValueError(f"[{entry_id}] sparse 词频为空，稀疏检索会整路失效")

        entry = record.get("entry") or {}
        for field_name in ("标题", "类别", "触发条件", "规则内容", "来源", "权威度", "核验状态"):
            if not entry.get(field_name):
                raise ValueError(f"[{entry_id}] entry.{field_name} 为空，注入 Prompt 时没内容")


# ---------------------------------------------------------------------------
# 检索参考实现
#
# 这不是产品代码（产品是 TypeScript，见设计文档 §5.4）。它存在的理由有两个：
#   1. 让「索引进去了但召不回」这类故障在构建期就暴露，而不是等界面做出来才发现
#   2. 作为 Plan 3 把检索搬到 TS 时的口径参照 —— 排序细节在这里是可执行、可核对的
# ---------------------------------------------------------------------------


def rank_dense(query: list[float], records: list[dict]) -> list[str]:
    scored = [
        (cosine(query, record["embedding"]), record["id"]) for record in records
    ]
    scored.sort(key=lambda item: (-item[0], item[1]))
    return [entry_id for _, entry_id in scored]


def rank_sparse(
    query_tf: dict[str, int], records: list[dict]
) -> list[str]:
    """BM25。文档侧词频来自构建期，df 在这里现算（语料就是全部条目，很小）。"""
    total_docs = len(records)
    doc_lengths = {
        record["id"]: sum(record["sparse"].values()) for record in records
    }
    avg_length = (sum(doc_lengths.values()) / total_docs) if total_docs else 0.0

    df: dict[str, int] = {}
    for record in records:
        for gram in record["sparse"]:
            df[gram] = df.get(gram, 0) + 1

    scored: list[tuple[float, str]] = []
    for record in records:
        score = 0.0
        length = doc_lengths[record["id"]]
        for gram, query_count in query_tf.items():
            tf = record["sparse"].get(gram)
            if not tf:
                continue
            doc_freq = df.get(gram, 0)
            idf = math.log(1 + (total_docs - doc_freq + 0.5) / (doc_freq + 0.5))
            denominator = tf + BM25_K1 * (
                1 - BM25_B + BM25_B * (length / avg_length if avg_length else 1.0)
            )
            score += idf * (tf * (BM25_K1 + 1) / denominator) * (1 + math.log(query_count))
        scored.append((score, record["id"]))

    scored.sort(key=lambda item: (-item[0], item[1]))
    return [entry_id for _, entry_id in scored]


def rrf_merge(rankings: list[list[str]], weights: list[float] | None = None) -> dict[str, float]:
    weights = weights or [1.0] * len(rankings)
    fused: dict[str, float] = {}
    for ranking, weight in zip(rankings, weights):
        for position, entry_id in enumerate(ranking):
            fused[entry_id] = fused.get(entry_id, 0.0) + weight / (RRF_K + position + 1)
    return fused


def retrieve(
    query_text: str,
    column_names: list[str],
    subject: str,
    records: list[dict],
    query_vector: list[float],
    top_k: int = TOP_K,
) -> list[tuple[str, float]]:
    """一次完整检索：学科过滤 → 双路召回 → RRF 融合 → 列名加权 → Top-K。"""
    candidates = [
        record
        for record in records
        if "通用" in record["学科"] or subject in record["学科"]
    ]
    if not candidates:
        return []

    dense = rank_dense(query_vector, candidates)
    sparse = rank_sparse(term_frequency(query_text), candidates)

    fused = rrf_merge([dense, sparse])

    # 列名精确命中 适用字段 → 加权。这是人工维护的强信号：
    # 「GB/T 2260」这类专有名词向量检索不敏感，而列名「行政区划」能精确命中。
    normalized_columns = {normalize_text(name).lower() for name in column_names}
    for record in candidates:
        entry_aliases = {alias.lower() for alias in record["适用字段"]}
        if entry_aliases and (entry_aliases & normalized_columns):
            fused[record["id"]] = fused.get(record["id"], 0.0) + FIELD_HIT_BOOST

    ordered = sorted(fused.items(), key=lambda item: (-item[1], item[0]))
    return ordered[:top_k]


def build_self_recall_cases(entries: list[Entry]) -> list[dict]:
    """机械生成自召回用例：由每条规则自己的别名与正例造一个「像真实数据摘要」的查询。

    诚实说明：这是**冒烟测试**，不是质量评测。它只能发现「索引坏了」（别名没进索引、
    学科过滤写反、向量没算），发现不了「规则本身写得不好」——
    后者需要团队标一份真数据集，见设计文档 §5.9 与 schema.md 第 6 节 Plan 4。
    """
    cases = []
    for entry in entries:
        columns = entry.适用字段[:3]
        values = entry.正例 or entry.标题
        query = " ".join(
            filter(None, ["列名: " + " ".join(columns) if columns else "", f"取值分布: {values}"])
        )
        cases.append(
            {
                "id": f"self-recall-{entry.id}",
                "query": query or entry.标题,
                "columns": columns,
                "subject": entry.适用学科,
                "expect": entry.id,
                "note": "机械生成，仅用于发现索引故障",
            }
        )
    return cases


def run_self_recall(
    entries: list[Entry],
    records: list[dict],
    embed_backend: str,
) -> dict:
    cases = build_self_recall_cases(entries)
    hits_at_1 = 0
    hits_at_k = 0
    reciprocal_ranks: list[float] = []
    failures: list[dict] = []

    for case in cases:
        if embed_backend == EMBEDDER_BGE:
            query_vector, _ = embed_all([case["query"]], embed_backend)
            vector = query_vector[0]
        else:
            vector = embed_hash_ngram(case["query"])

        ranked = retrieve(
            case["query"], case["columns"], case["subject"], records, vector
        )
        ids = [entry_id for entry_id, _ in ranked]

        if ids and ids[0] == case["expect"]:
            hits_at_1 += 1
        if case["expect"] in ids:
            hits_at_k += 1
            reciprocal_ranks.append(1.0 / (ids.index(case["expect"]) + 1))
        else:
            reciprocal_ranks.append(0.0)
            failures.append(
                {
                    "case": case["id"],
                    "query": case["query"],
                    "expect": case["expect"],
                    "got": ids,
                }
            )

    total = len(cases) or 1
    return {
        "总用例": len(cases),
        "Hit@1": round(hits_at_1 / total, 4),
        "Hit@%d" % TOP_K: round(hits_at_k / total, 4),
        "MRR": round(sum(reciprocal_ranks) / total, 4),
        "未命中明细": failures,
        "说明": (
            "机械生成的自召回冒烟测试，用于发现索引故障（别名未进索引、"
            "学科过滤写反、向量未算）。不是检索质量评测 —— "
            "真实评测需要团队标注的数据摘要→规则对，见设计文档 §5.9。"
        ),
    }


# ---------------------------------------------------------------------------
# 字段别名标签体系（两级）
#
# 「适用字段」是检索最可靠的信号，但同一个字段在不同规则里写法不同。
# 这里把别名按「共同出现在同一条规则里」聚成组（连通分量），
# 产出一份可追溯的归并结果：每个组都记着是哪几条规则把它连起来的。
#
# 刻意不做同义判定 —— 「学历」和「文化程度」是不是同一个字段，是领域判断，
# 该由填表的人决定。脚本只做「你们已经写在一起了，那它们就是一组」这件事。
# ---------------------------------------------------------------------------


def build_alias_taxonomy(entries: list[Entry]) -> dict:
    """把 适用字段 的别名聚成两级标签体系。

    聚类依据只有一条：**共同出现在同一条规则里**。脚本不做任何同义判定 ——
    「学历」和「文化程度」是不是同一个字段是领域判断，该由填表的人决定。
    脚本只回答「你们已经把它们写在一起了，那它们就是一组」，
    并记下是哪几条规则把它们连起来的，让结果可逐组回查。
    """
    parent: dict[str, str] = {}
    first_seen: dict[str, int] = {}
    evidence: dict[str, list[str]] = {}
    order = 0

    def find(node: str) -> str:
        parent.setdefault(node, node)
        while parent[node] != node:
            parent[node] = parent[parent[node]]
            node = parent[node]
        return node

    def union(left: str, right: str) -> None:
        root_left, root_right = find(left), find(right)
        if root_left != root_right:
            parent[root_right] = root_left

    for entry in entries:
        aliases = entry.适用字段
        for alias in aliases:
            find(alias)
            if alias not in first_seen:
                first_seen[alias] = order
                order += 1
            evidence.setdefault(alias, []).append(entry.id)
        for alias in aliases[1:]:
            union(aliases[0], alias)

    groups: dict[str, list[str]] = {}
    for alias in parent:
        groups.setdefault(find(alias), []).append(alias)

    alias_subject: dict[str, str] = {}
    alias_count: dict[str, int] = {}
    for entry in entries:
        alias_count[entry.id] = len(entry.适用字段)
        for alias in entry.适用字段:
            alias_subject.setdefault(alias, entry.适用学科)

    tree: dict[str, list[dict]] = {}
    for members in groups.values():
        # 组名取「在 Excel 里最早出现的那个别名」——填表人通常把规范写法写在最前，
        # 按字母序取会得到 gender／剂量 这种半中半英的怪标签
        ordered = sorted(members, key=lambda alias: first_seen.get(alias, 1 << 30))
        subject = alias_subject.get(ordered[0], "通用")
        sources = sorted({eid for alias in members for eid in evidence.get(alias, [])})

        group = {
            "别名组": ordered[0],
            "别名词表": ordered,
            "规模": len(members),
            "依据": sources,
        }

        # 单条规则一口气声明 5 个以上别名 → 列进「请团队判断」的候选。
        #
        # 刻意只提出疑问，不下「它管了多件事」的结论：`学号／编号／工号／id／样本号`
        # 是同族的标识列，`性别／学历／民族／行政区划` 不是 —— 靠字符 2-gram 分不出这两者
        # （`学号` 与 `编号` 的字面重合度是 0），需要领域判断。
        # 所以这里产出的是**候选人清单，不是缺陷列表**：别名多本身是对的（别名就是给检索用的），
        # 只有它们分属不同字段族时才该拆。判定权交回填表的人。
        widest = max(sources, key=lambda eid: alias_count.get(eid, 0), default=None)
        width = alias_count.get(widest, 0) if widest else 0
        if width >= 5:
            group["待确认"] = (
                f"该组 {len(members)} 个别名里有 {width} 个来自同一条规则 `{widest}`。"
                f"请判断这 {width} 个是否属于同一字段族：是则该条管住即可，"
                f"若跨了不同字段族（如编码体系互不相干的几类字段）则建议拆开。"
            )

        tree.setdefault(subject, []).append(group)

    for subject in tree:
        tree[subject].sort(key=lambda item: (-item["规模"], item["别名组"]))

    generalized = [entry.id for entry in entries if not entry.适用字段]

    return {
        "说明": (
            "两级标签体系：一级=适用学科，二级=字段别名组。"
            "别名组由「共同出现在同一条规则里」聚类得出，依据字段列出是哪几条规则"
            "把它们连起来的 —— 可逐组回查，脚本未做任何同义判定。"
        ),
        "体系": tree,
        "泛化规则": {
            "说明": "适用字段留空 = 任何列都适用，不参与列名加权（schema.md 第 2 节）",
            "条目": generalized,
        },
    }


# ---------------------------------------------------------------------------
# 入库质量分
#
# 分数是**可解释的扣分制**，每一项都能追回 Excel 的具体单元格。
# 不给「综合评分」这类黑箱数字 —— 团队要能知道改哪里能把分提上去。
# ---------------------------------------------------------------------------

QUALITY_ITEMS = [
    ("必填字段齐全", 20, "id／标题／触发条件／规则内容／来源 五项齐全"),
    ("来源可追溯", 15, "来源不是「待确认」占位，能追到国标号／口径名／实验记录"),
    ("核验状态已核对", 10, "团队对照原文核过。待核对的条目等于「内容还没人验过」，不能算入库质量达标"),
    ("反例齐备", 15, "归并类规则必须有反例；非归并类给反例也加分"),
    ("别名充分", 15, "适用字段 ≥ 3 个别名，或明确是泛化规则（留空）"),
    ("单条一事", 15, "规则内容 2–5 句，且标题不超过 30 字"),
    ("正例齐备", 10, "给了「对的样子」，比只讲规则更有效"),
]


def score_entry(entry: Entry) -> dict:
    score = 0
    breakdown: list[dict] = []
    got: dict[str, bool] = {}

    required_ok = all(
        getattr(entry, name) for name in ("id", "标题", "触发条件", "规则内容", "来源")
    )
    got["必填字段齐全"] = bool(required_ok)

    got["来源可追溯"] = bool(entry.来源) and entry.来源 not in ("待确认", "待核对")
    got["核验状态已核对"] = entry.核验状态 == "已核对"
    got["反例齐备"] = bool(entry.反例)
    got["别名充分"] = len(entry.适用字段) >= 3 or not entry.适用字段

    sentences = len([s for s in re.split(r"[。！？；;]", entry.规则内容) if s.strip()])
    got["单条一事"] = bool(entry.规则内容) and 2 <= sentences <= 5 and len(entry.标题) <= 30

    got["正例齐备"] = bool(entry.正例)

    for name, weight, why in QUALITY_ITEMS:
        passed = got[name]
        if passed:
            score += weight
        breakdown.append(
            {"项": name, "权重": weight, "得分": weight if passed else 0, "判据": why}
        )

    return {"id": entry.id, "标题": entry.标题, "入库质量分": score, "明细": breakdown}


def percentile(sorted_values: list[float], fraction: float) -> float:
    """分位数。输入须已升序。空列表返回 0。"""
    if not sorted_values:
        return 0.0
    position = min(len(sorted_values) - 1, max(0, int(round(fraction * (len(sorted_values) - 1)))))
    return sorted_values[position]


def analyze_similarity(
    entries: list[Entry],
    embeddings: list[list[float]],
    threshold: float,
    top_n: int = SIMILARITY_TOP_N,
) -> dict:
    """同科条目两两余弦：算分布、筛疑似近重复、并固定给出最相似的前几名。

    **无论有没有过阈值都要给出分布与 TOP 列表。** 一个常年报 0 的指标是死的 ——
    真出过一次：哈希后端的余弦最大只有 0.38，而阈值写死在 0.86，于是这个检查永远沉默。
    把分布摊开会话，阈值是否校准就一眼可见。
    """
    pairs: list[dict] = []
    for i in range(len(entries)):
        for j in range(i + 1, len(entries)):
            if entries[i].适用学科 != entries[j].适用学科:
                continue
            pairs.append(
                {
                    "a": entries[i].id,
                    "b": entries[j].id,
                    "余弦": round(cosine(embeddings[i], embeddings[j]), 4),
                }
            )

    pairs.sort(key=lambda item: (-item["余弦"], item["a"], item["b"]))
    values = sorted(item["余弦"] for item in pairs)

    return {
        "阈值": threshold,
        "分布": {
            "对数": len(pairs),
            "最大": round(values[-1], 4) if values else 0.0,
            "p99": round(percentile(values, 0.99), 4),
            "p90": round(percentile(values, 0.90), 4),
            "中位": round(percentile(values, 0.50), 4),
        },
        "疑似近重复": [item for item in pairs if item["余弦"] >= threshold],
        "最相似": pairs[:top_n],
    }


# ---------------------------------------------------------------------------
# 报告
# ---------------------------------------------------------------------------


def render_report(
    entries: list[Entry],
    diagnostics: list[Diagnostic],
    scores: list[dict],
    similarity: dict,
    taxonomy: dict,
    rules_meta: dict,
    self_recall: dict | None,
) -> str:
    errors = [item for item in diagnostics if item.level == "error"]
    warnings = [item for item in diagnostics if item.level == "warning"]
    pending = [entry for entry in entries if entry.核验状态 == "待核对"]

    by_sheet: dict[str, int] = {}
    by_subject: dict[str, int] = {}
    by_authority: dict[str, int] = {}
    for entry in entries:
        by_sheet[entry.类别] = by_sheet.get(entry.类别, 0) + 1
        by_subject[entry.适用学科] = by_subject.get(entry.适用学科, 0) + 1
        by_authority[entry.权威度] = by_authority.get(entry.权威度, 0) + 1

    lines: list[str] = []
    add = lines.append

    add("# 知识库构建报告")
    add("")
    add(f"> 由 `knowledge/tools/build-index.py` 生成。真源是 `knowledge/知识库模板.xlsx`。")
    add("")
    add("## 一、结论")
    add("")
    add("| 项 | 值 |")
    add("| --- | --- |")
    add(f"| 条目总数 | {len(entries)} |")
    add(f"| 错误 | {len(errors)} |")
    add(f"| 警告 | {len(warnings)} |")
    add(f"| 待核对条目 | {len(pending)} |")
    add(f"| 疑似近重复对 | {len(similarity['疑似近重复'])}（阈值 {similarity['阈值']}，实际最大 {similarity['分布']['最大']}） |")
    add(f"| 向量后端 | {rules_meta['model']} |")
    add(f"| 构建时间 | {rules_meta['builtAt']} |")
    if self_recall:
        add(f"| 自召回 Hit@{rules_meta['topK']} | {self_recall['Hit@%d' % rules_meta['topK']]} |")
        add(f"| 自召回 MRR | {self_recall['MRR']} |")
    add("")

    add("## 二、分布")
    add("")
    add("| 维度 | 分布 |")
    add("| --- | --- |")
    add("| 类别 | " + "　".join(f"{k} {v}" for k, v in by_sheet.items()) + " |")
    add("| 学科 | " + "　".join(f"{k} {v}" for k, v in by_subject.items()) + " |")
    add("| 权威度 | " + "　".join(f"{k} {v}" for k, v in by_authority.items()) + " |")
    add("")

    add("## 三、错误")
    add("")
    if errors:
        add("构建已中止，下面的产物不会被写出。")
        add("")
        for item in errors:
            add(f"- {item.format()}")
    else:
        add("（无）")
    add("")

    add("## 四、警告")
    add("")
    if warnings:
        for item in warnings:
            add(f"- {item.format()}")
    else:
        add("（无）")
    add("")

    add("## 五、待核对清单")
    add("")
    add("这些条目 `核验状态 = 待核对`。答辩要能逐条追溯出处，请对照原文核对后改成「已核对」。")
    add("")
    if pending:
        add("| id | 标题 | 来源 |")
        add("| --- | --- | --- |")
        for entry in pending:
            add(f"| `{entry.id}` | {entry.标题} | {entry.来源} |")
    else:
        add("（无）")
    add("")

    add("## 六、入库质量分（升序，先看最后几行）")
    add("")
    add("扣分项写给团队看：改哪里能把分提上去。权重定义见 `build-index.py` 的 `QUALITY_ITEMS`。")
    add("")
    add("| 分数 | id | 标题 | 未得分项 |")
    add("| --- | --- | --- | --- |")
    for item in sorted(scores, key=lambda row: (row["入库质量分"], row["id"])):
        missed = "、".join(
            step["项"] for step in item["明细"] if step["得分"] == 0
        )
        add(f"| {item['入库质量分']} | `{item['id']}` | {item['标题']} | {missed or '—'} |")
    add("")

    add("## 七、字段别名体系")
    add("")
    add(
        "一级=学科，二级=字段别名组。别名组由「共同出现在同一条规则里」聚类得出 —— "
        "脚本不做同义判定，只把「你们已经写在一起的」归成一组，并记下依据。"
    )
    add("")
    for subject, groups in taxonomy["体系"].items():
        add(f"### {subject}")
        add("")
        add("| 别名组 | 规模 | 词表 | 依据 |")
        add("| --- | --- | --- | --- |")
        for group in groups:
            add(
                f"| {group['别名组']} | {group['规模']} | "
                f"{'、'.join(group['别名词表'])} | {'、'.join('`' + e + '`' for e in group['依据'])} |"
            )
        add("")
        for group in groups:
            if "待确认" in group:
                add(f"- **待确认 · {group['别名组']}**：{group['待确认']}")
        if any("待确认" in group for group in groups):
            add("")
            add(
                "> 上面这份是**候选人清单，不是缺陷列表**。别名多本身是对的 —— 别名就是给检索用的，"
                "写全才能命中不同表头。只有当一个字段族的别名里混进了另一个字段族（如"
                "`学号／工号／编号` 与 `性别／学历／民族` 混在一条里），才该拆。"
                "机器分不出这两者（`学号` 与 `编号` 的字面重合度是 0），所以判定权交回填表的人。"
            )
        add("")

    generalized = taxonomy["泛化规则"]["条目"]
    add(f"泛化规则（适用字段留空，任何列都适用，不参与列名加权）：{len(generalized)} 条")
    add("")
    if generalized:
        for entry_id in generalized:
            add(f"- `{entry_id}`")
        add("")

    add("## 八、同科条目相似度")
    add("")
    add(
        "同一学科内两两比较的余弦相似度。**为什么连分布一起报**："
        "近重复阈值若够不着，这个检查就会长期沉默，看起来像「没有问题」。"
        "把分布摊开，阈值是否校准一眼可见。"
    )
    add("")
    add(
        "> **阈值精度有限，请当提示看而不是判决看。** 当前离线后端的余弦尺度是被压扁的，"
        "上面这个阈值只能照实测分布定，**没有标注集可校准**，因此既有漏报也有误报。"
        "换 bge 语义后端后阈值必须重定，Plan 4 建起评测集才谈得上校准。"
    )
    add("")
    distribution = similarity["分布"]
    add("| 对数 | 最大 | p99 | p90 | 中位 | 本条阈值 |")
    add("| --- | --- | --- | --- | --- | --- |")
    add(
        f"| {distribution['对数']} | {distribution['最大']} | {distribution['p99']} | "
        f"{distribution['p90']} | {distribution['中位']} | {similarity['阈值']} |"
    )
    add("")
    add(f"**最相似 TOP {len(similarity['最相似'])}**（无条件展示，供人工判断）")
    add("")
    add("| 余弦 | A | B |")
    add("| --- | --- | --- |")
    for pair in similarity["最相似"]:
        add(f"| {pair['余弦']} | `{pair['a']}` | `{pair['b']}` |")
    add("")
    if similarity["疑似近重复"]:
        add(f"**超过阈值 {similarity['阈值']} 的 {len(similarity['疑似近重复'])} 对**")
        add("")
        add("| 余弦 | A | B |")
        add("| --- | --- | --- |")
        for pair in similarity["疑似近重复"]:
            add(f"| {pair['余弦']} | `{pair['a']}` | `{pair['b']}` |")
        add("")
        add("两条可能在讲同一件事。确认后请删掉一条，或把边界写进反例。")
        add("")
    else:
        add(f"**没有一对超过阈值 {similarity['阈值']}。**")
        add("")
        if distribution["最大"] < similarity["阈值"]:
            add(
                f"> ⚠️ 阈值够不着：最大值 {distribution['最大']} 小于阈值 {similarity['阈值']}。"
                f"说明这个后端下该指标无法生效 —— 请按上面的分布重定阈值。"
                f"（语义模型与哈希后端的余弦尺度差很多，阈值必须与后端匹配）"
            )
        else:
            add(f"> 最大值 {distribution['最大']} 已接近但未超过阈值，属正常。")
        add("")

    if self_recall:
        add("## 九、自召回冒烟测试")
        add("")
        add(f"> {self_recall['说明']}")
        add("")
        add("| 指标 | 值 |")
        add("| --- | --- |")
        add(f"| 总用例 | {self_recall['总用例']} |")
        add(f"| Hit@1 | {self_recall['Hit@1']} |")
        add(f"| Hit@{rules_meta['topK']} | {self_recall['Hit@%d' % rules_meta['topK']]} |")
        add(f"| MRR | {self_recall['MRR']} |")
        add("")
        if self_recall["未命中明细"]:
            add("### 未命中明细")
            add("")
            for failure in self_recall["未命中明细"]:
                add(f"- `{failure['expect']}` 未进 Top-{rules_meta['topK']}")
                add(f"  - 查询：{failure['query']}")
                add(f"  - 实际召回：{'、'.join(failure['got']) or '（空）'}")
            add("")

    add("---")
    add("")
    add("本报告由 AI 依据 `knowledge/schema.md` 的规格生成。分数、指标均可由构建脚本复现。")
    add("")

    return "\n".join(lines)


# ---------------------------------------------------------------------------
# 自测
#
# 项目约定是「每个模块都有先失败后通过的测试」。这是 Python 工具链，
# 跑不进 Vitest，所以测试内嵌成 `--self-test`，一条命令就能跑：
# 不依赖真实模板、不联网、几百毫秒完事。改动本脚本后必须跑它。
# ---------------------------------------------------------------------------

VALID_ROW = [
    "general-demo-01",
    "示例规则",
    "通用",
    "学历,受教育程度,edu",
    "该列出现不一致写法时",
    "两条规则。用于自测，内容本身没有意义。",
    "硕士 ← 研究生",
    "硕士 与 博士 不可归并",
    "自测构造",
    "行业",
    "待核对",
]


class Checker:
    def __init__(self) -> None:
        self.passed = 0
        self.failures: list[str] = []

    def ok(self, condition: bool, label: str) -> None:
        if condition:
            self.passed += 1
        else:
            self.failures.append(label)


def _write_workbook(path: Path, sheets: dict[str, list[list]]) -> None:
    from openpyxl import Workbook

    book = Workbook()
    book.remove(book.active)
    for name, rows in sheets.items():
        sheet = book.create_sheet(name)
        sheet.append(COLUMNS)
        for row in rows:
            sheet.append(row)
    book.save(path)


def _diagnose(workbook_path: Path) -> list[Diagnostic]:
    diagnostics: list[Diagnostic] = []
    entries = load_entries(workbook_path, diagnostics)
    validate(entries, diagnostics)
    return diagnostics


def _has_error(diagnostics: list[Diagnostic], field: str, fragment: str = "") -> bool:
    return any(
        item.level == "error"
        and item.field == field
        and (fragment in item.message)
        for item in diagnostics
    )


def _has_warning(diagnostics: list[Diagnostic], field: str, fragment: str = "") -> bool:
    return any(
        item.level == "warning"
        and item.field == field
        and (fragment in item.message)
        for item in diagnostics
    )


def run_self_test() -> int:
    import tempfile

    check = Checker()

    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)

        # --- 正向：合法工作簿不报错 ---
        good = root / "good.xlsx"
        _write_workbook(
            good,
            {
                "数据标准": [VALID_ROW],
                "口径规则": [],
                "清洗案例": [],
                "填写说明": [["便签内容"], ["不该被当成条目"]],
                "字段说明": [["列名", "怎么填"]],
            },
        )
        diagnostics = _diagnose(good)
        check.ok(
            not any(item.level == "error" for item in diagnostics),
            "合法工作簿不应报错误，实际：" + "; ".join(d.message for d in diagnostics if d.level == "error"),
        )

        # --- 文档性 sheet 不被当成条目 ---
        entries = load_entries(good, [])
        check.ok(len(entries) == 1, f"只应产出 1 条条目，实际 {len(entries)}")
        check.ok(
            entries[0].类别 == "数据标准",
            f"类别应由 sheet 名推断，实际 {entries[0].类别}",
        )
        check.ok(
            entries[0].适用字段 == ["学历", "受教育程度", "edu"],
            f"逗号分隔的别名应拆成列表，实际 {entries[0].适用字段}",
        )

        # --- 反向：各类错误都能被拦下 ---
        def one_row(column: int, value) -> list[Diagnostic]:
            row = list(VALID_ROW)
            row[column] = value
            path = root / f"bad-{column}-{str(value)[:8]}.xlsx"
            _write_workbook(path, {"数据标准": [row], "口径规则": [], "清洗案例": []})
            return _diagnose(path)

        check.ok(_has_error(one_row(1, ""), "标题"), "空标题应报错")
        # 注意列序：0=id 1=标题 2=适用学科 3=适用字段 4=触发条件 5=规则内容
        #            6=正例 7=反例 8=来源 9=权威度 10=核验状态
        check.ok(_has_error(one_row(8, ""), "来源", "为空"), "空来源应报错")
        check.ok(_has_error(one_row(5, "同上"), "规则内容"), "规则内容含「同上」应报错")
        check.ok(
            _has_error(one_row(9, "很权威"), "权威度"), "权威度不在枚举内应报错"
        )
        check.ok(
            _has_error(one_row(10, "待审核"), "核验状态"), "核验状态不在枚举内应报错"
        )
        check.ok(
            _has_error(one_row(0, "General-Demo"), "id"), "id 含非法字符应报错"
        )

        # --- 权威度与来源的对应：旧清单之外，URL / GB 编号也算可核对证据 ---
        url_row = list(VALID_ROW)
        url_row[8] = (
            "pandas 官方文档：read_csv（读取参数）\n"
            "https://pandas.pydata.org/docs/reference/api/pandas.read_csv.html\n"
            "核对日期：2026-10-25"
        )
        url_row[9] = "行业"
        path = root / "authority-url.xlsx"
        _write_workbook(path, {"数据标准": [url_row], "口径规则": [], "清洗案例": []})
        check.ok(
            not _has_warning(_diagnose(path), "权威度"),
            "行业来源带 URL 时应视为可核对依据，不再报「不在清单」警告",
        )

        gb_row = list(VALID_ROW)
        gb_row[8] = "国家标准信息公共服务平台：GB/T 4658-2006（学历代码）"
        gb_row[9] = "国标"
        path = root / "authority-gb.xlsx"
        _write_workbook(path, {"数据标准": [gb_row], "口径规则": [], "清洗案例": []})
        check.ok(
            not _has_warning(_diagnose(path), "权威度"),
            "国标来源带 GB/T 编号时应可核对，不再报「不在清单」警告",
        )

        bare_row = list(VALID_ROW)
        bare_row[8] = "我们团队自己定的口径"
        bare_row[9] = "行业"
        path = root / "authority-bare.xlsx"
        _write_workbook(path, {"数据标准": [bare_row], "口径规则": [], "清洗案例": []})
        check.ok(
            _has_warning(_diagnose(path), "权威度"),
            "行业来源既不在清单、也没有 URL 时仍要报警告（收窄不能滑成放水）",
        )

        # 归并类规则缺反例
        row = list(VALID_ROW)
        row[1] = "硕士与研究生应归并"
        row[7] = ""
        path = root / "merge-no-counter.xlsx"
        _write_workbook(path, {"数据标准": [row], "口径规则": [], "清洗案例": []})
        check.ok(_has_error(_diagnose(path), "反例"), "归并类规则缺反例应报错")

        # id 重复
        duplicate = list(VALID_ROW)
        duplicate[1] = "另一条规则"
        path = root / "dup-id.xlsx"
        _write_workbook(path, {"数据标准": [VALID_ROW, duplicate], "口径规则": [], "清洗案例": []})
        check.ok(_has_error(_diagnose(path), "id", "重复"), "id 重复应报错")

        # 「略」的窄化：正常用词不该被当成占位写法（防误报）
        for order, normal_text in enumerate(
            (
                "策略上应忽略短期波动。这是第二句。",
                "必须先把哨兵值规范成缺失，再决定填补策略。这是第二句。",
                "这一条可以省略。这是第二句。",
            )
        ):
            row = list(VALID_ROW)
            row[5] = normal_text
            path = root / f"normal-{order}.xlsx"
            _write_workbook(path, {"数据标准": [row], "口径规则": [], "清洗案例": []})
            check.ok(
                not _has_error(_diagnose(path), "规则内容"),
                f"正常用词「策略／略过／省略」不应被判成占位写法：{normal_text}",
            )

        # 反过来：真正的占位写法必须被抓到，否则收窄就成了放水
        for order, placeholder in enumerate(("前半句已述，此处略。", "（略）", "同上。")):
            row = list(VALID_ROW)
            row[5] = placeholder
            path = root / f"placeholder-{order}.xlsx"
            _write_workbook(path, {"数据标准": [row], "口径规则": [], "清洗案例": []})
            check.ok(
                _has_error(_diagnose(path), "规则内容"),
                f"真正的占位写法应被抓到：{placeholder}",
            )

        # 未知 sheet
        path = root / "unknown-sheet.xlsx"
        _write_workbook(
            path,
            {"数据标准": [VALID_ROW], "口径规则": [], "清洗案例": [], "数据字典": []},
        )
        check.ok(_has_error(_diagnose(path), "sheet 名"), "未知 sheet 名应报错")

        # --- 别名冲突分「同类别内」与「跨类别」两类 ---
        def alias_row(entry_id: str, aliases: str) -> list:
            row = list(VALID_ROW)
            row[0] = entry_id
            row[3] = aliases
            return row

        path = root / "alias-same.xlsx"
        _write_workbook(
            path,
            {
                "数据标准": [alias_row("general-demo-a", "学历,文化程度"), alias_row("general-demo-b", "学历,edu")],
                "口径规则": [],
                "清洗案例": [],
            },
        )
        diagnostics = _diagnose(path)
        check.ok(
            any(item.level == "warning" and "同属「数据标准」" in item.message for item in diagnostics),
            "同类别内共用别名应报「同属」警告",
        )

        path = root / "alias-cross.xlsx"
        _write_workbook(
            path,
            {
                "数据标准": [alias_row("general-demo-a", "学历,文化程度")],
                "口径规则": [alias_row("general-demo-c", "学历,文化程度")],
                "清洗案例": [],
            },
        )
        diagnostics = _diagnose(path)
        check.ok(
            not any("同属" in item.message for item in diagnostics),
            "跨类别共用别名不应报同类别警告（那是约定的分工）",
        )
        check.ok(
            any("跨类别的多条规则共用" in item.message for item in diagnostics),
            "跨类别共用别名应汇总成一行，不能逐个别名刷警告",
        )

        # 未登记学科
        row = list(VALID_ROW)
        row[2] = "考古学"
        path = root / "unknown-subject.xlsx"
        _write_workbook(path, {"数据标准": [row], "口径规则": [], "清洗案例": []})
        check.ok(_has_error(_diagnose(path), "适用学科"), "未登记 slug 的学科应报错")

        # 表头被改
        from openpyxl import Workbook

        path = root / "bad-header.xlsx"
        book = Workbook()
        book.remove(book.active)
        sheet = book.create_sheet("数据标准")
        sheet.append(["id", "名称"] + COLUMNS[2:])
        book.create_sheet("口径规则").append(COLUMNS)
        book.create_sheet("清洗案例").append(COLUMNS)
        book.save(path)
        check.ok(_has_error(_diagnose(path), "表头"), "表头与约定不符应报错")

        # --- 空行跳过 ---
        path = root / "blank-rows.xlsx"
        _write_workbook(
            path,
            {
                "数据标准": [VALID_ROW] + [[None] * len(COLUMNS)] * 40,
                "口径规则": [],
                "清洗案例": [],
            },
        )
        entries = load_entries(path, [])
        check.ok(len(entries) == 1, f"40 个空行不应产出条目，实际 {len(entries)} 条")

        # --- YAML 往返保内容 ---
        entries = load_entries(good, [])
        yaml_text = dump_entries_yaml(entries, "通用", "general")
        reloaded = require_yaml().safe_load(yaml_text)
        check.ok(len(reloaded) == 1, "YAML 重载后条目数应一致")
        for key, value in entries[0].to_yaml_dict().items():
            check.ok(
                reloaded[0][key] == value,
                f"YAML 往返改变了字段 {key}：{reloaded[0][key]!r} != {value!r}",
            )

        # --- 向量：确定性 + 归一化 ---
        vector_a = embed_hash_ngram("缺失编码的判定")
        vector_b = embed_hash_ngram("缺失编码的判定")
        check.ok(vector_a == vector_b, "同一文本两次向量化结果应一致（跨进程也须一致）")
        norm = math.sqrt(sum(value * value for value in vector_a))
        check.ok(abs(norm - 1.0) < 1e-9, f"向量应 L2 归一化，实际模长 {norm}")
        check.ok(
            abs(cosine(vector_a, embed_hash_ngram("缺失编码的判定"))) > 0.999,
            "同一文本的余弦相似度应为 1",
        )

        # --- 2-gram 跨空白 ---
        grams_with_space = term_frequency("硕士 研究生")
        grams_without = term_frequency("硕士研究生")
        check.ok(
            "士研" in grams_with_space and "士研" in grams_without,
            "去空白后再切 gram，带空格的写法与不带空格应共享 gram",
        )

        # --- 相似度分析：阈值够不着时也必须给出分布与 TOP ---
        check.ok(percentile([], 0.99) == 0.0, "空列表的分位数应为 0")
        check.ok(percentile([1.0, 2.0, 3.0], 0.0) == 1.0, "分位数下界应取下端")
        check.ok(percentile([1.0, 2.0, 3.0], 1.0) == 3.0, "分位数上界应取上端")

        # 造一对同科条目：只有 1 对，便于断言
        pair = [
            Entry(
                id="general-demo-a", 标题="缺失编码", 适用学科="通用",
                规则内容="负数编码表示未回答。这是第二句。", 触发条件="自测", 来源="自测构造",
            ),
            Entry(
                id="general-demo-b", 标题="缺失编码的判定", 适用学科="通用",
                规则内容="负数编码表示未回答的原因。这是第二句。", 触发条件="自测", 来源="自测构造",
            ),
        ]
        pair_vectors, _ = embed_all([item.dense_text() for item in pair], EMBEDDER_HASH)

        # 阈值故意设成 0.999：一对都过不了，但分布与 TOP 必须照旧给出
        hopeless = analyze_similarity(pair, pair_vectors, threshold=0.999)
        check.ok(hopeless["疑似近重复"] == [], "阈值极高时不应有疑似近重复")
        check.ok(
            len(hopeless["最相似"]) == 1,
            "阈值够不着时仍须给出最相似列表，否则无法区分「没重复」与「阈值失效」",
        )
        check.ok(
            hopeless["分布"]["对数"] == 1 and "p99" in hopeless["分布"],
            "分布统计必须存在，否则阈值是否校准无从判断",
        )

        # 阈值放到 0 时那一对必须被标出来（自检要有牙齿，不能永远返回空）
        flagged = analyze_similarity(pair, pair_vectors, threshold=0.0)
        check.ok(
            [item["a"] for item in flagged["疑似近重复"]] == ["general-demo-a"],
            "阈值设为 0 时唯一那一对必须被标为疑似近重复",
        )

        # --- rules.json：向量压成一行、字符串数组保持多行、回读自检有效 ---
        demo = load_entries(good, [])
        vectors, model = embed_all([item.dense_text() for item in demo], EMBEDDER_HASH)
        payload = build_rules_json(demo, vectors, model, "2026-01-01T00:00:00Z")
        rules_path = root / "rules.json"
        write_json(rules_path, payload, compact_number_arrays=True)
        rules_text = rules_path.read_text(encoding="utf-8")

        check.ok(
            len(rules_text.split("\n")) < 120,
            f"512 维向量应压成一行，实际 {len(rules_text.splitlines())} 行",
        )
        check.ok(
            '"受教育程度",\n' in rules_text,
            "含字符串的数组（适用字段）应保持多行，别名要能一眼看清",
        )

        try:
            verify_rules_json(rules_path, len(demo), EMBEDDING_DIM)
            check.ok(True, "rules.json 回读自检应通过")
        except Exception as exc:  # noqa: BLE001 - 自测就是要把失败报出来
            check.ok(False, f"rules.json 回读自检不应失败：{exc}")

        # 自检要有牙齿：篡改产物后必须报错，否则它只是走过场
        tampered = root / "tampered.json"
        bad = json.loads(rules_text)
        bad["entries"][0]["embedding"] = bad["entries"][0]["embedding"][:16]
        tampered.write_text(json.dumps(bad, ensure_ascii=False), encoding="utf-8")
        try:
            verify_rules_json(tampered, len(demo), EMBEDDING_DIM)
            check.ok(False, "向量维度被截断时应报错，但回读自检放过了")
        except ValueError:
            check.ok(True, "向量维度被截断时应报错")

        # 规则正文里的字面量反斜杠（general-fullwidth 有 \r）要能原样穿过 JSON
        tricky = Entry(
            id="general-demo-escape",
            标题="转义测试",
            规则内容="行尾多余的 \\r 与引号「\"」都要原样保留。这是第二句。",
            触发条件="自测",
            来源="自测构造",
        )
        tricky_vectors, tricky_model = embed_all([tricky.dense_text()], EMBEDDER_HASH)
        tricky_path = root / "escape.json"
        write_json(
            tricky_path,
            build_rules_json([tricky], tricky_vectors, tricky_model, "2026-01-01T00:00:00Z"),
            compact_number_arrays=True,
        )
        round_tripped = json.loads(tricky_path.read_text(encoding="utf-8"))
        check.ok(
            round_tripped["entries"][0]["entry"]["规则内容"] == tricky.规则内容,
            f"规则内容里的反斜杠与引号应原样往返，实际 {round_tripped['entries'][0]['entry']['规则内容']!r}",
        )

    print(f"自测通过 {check.passed} 项")
    if check.failures:
        print(f"失败 {len(check.failures)} 项：", file=sys.stderr)
        for failure in check.failures:
            print(f"  ✗ {failure}", file=sys.stderr)
        return 1
    return 0


# ---------------------------------------------------------------------------
# 主流程
# ---------------------------------------------------------------------------


def resolve_built_at() -> str:
    """可复现构建：设了 SOURCE_DATE_EPOCH 就用它，否则用当前时间。"""
    epoch = os.environ.get("SOURCE_DATE_EPOCH")
    if epoch and epoch.isdigit():
        moment = datetime.fromtimestamp(int(epoch), tz=timezone.utc)
    else:
        moment = datetime.now(timezone.utc)
    return moment.replace(microsecond=0).isoformat().replace("+00:00", "Z")


def parse_args(argv: list[str]) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="knowledge/知识库模板.xlsx → entries/*.yaml → index/rules.json"
    )
    parser.add_argument("--template", type=Path, default=TEMPLATE)
    parser.add_argument("--check", action="store_true", help="只校验，不写任何文件")
    parser.add_argument(
        "--self-recall",
        action="store_true",
        help="构建后跑自召回冒烟测试（会写进报告）",
    )
    parser.add_argument(
        "--embedder",
        choices=[EMBEDDER_HASH, EMBEDDER_BGE],
        default=EMBEDDER_HASH,
        help="dense 向量后端。默认离线哈希后端；语义模型需装 sentence-transformers",
    )
    parser.add_argument(
        "--self-test",
        action="store_true",
        help="跑内嵌的自测（不读真实模板，不写文件）",
    )
    return parser.parse_args(argv)


def main(argv: list[str]) -> int:
    args = parse_args(argv)

    if args.self_test:
        return run_self_test()

    diagnostics: list[Diagnostic] = []

    entries = load_entries(args.template, diagnostics)
    validate(entries, diagnostics)

    errors = [item for item in diagnostics if item.level == "error"]
    if errors:
        print("构建中止：校验未通过。", file=sys.stderr)
        for item in errors:
            print(f"  ✗ {item.format()}", file=sys.stderr)
        print(
            f"\n共 {len(errors)} 个错误。请改 Excel 后重跑；本次没有写出任何文件。",
            file=sys.stderr,
        )
        return 1

    warnings = [item for item in diagnostics if item.level == "warning"]

    if not entries:
        print("模板里一条规则都没有，无事可做。", file=sys.stderr)
        return 1

    if args.check:
        print(f"校验通过：{len(entries)} 条，{len(warnings)} 条警告。")
        for item in warnings:
            print(f"  ! {item.format()}")
        return 0

    # --- 向量 ---
    embeddings, model = embed_all([entry.dense_text() for entry in entries], args.embedder)

    built_at = resolve_built_at()
    rules = build_rules_json(entries, embeddings, model, built_at)

    # --- 写 entries/<学科>.yaml ---
    ENTRIES_DIR.mkdir(parents=True, exist_ok=True)
    # 先清掉上一次的产物：学科改名后旧文件会变成孤儿，留在那里会被误读成现行规则
    for stale in ENTRIES_DIR.glob("*.yaml"):
        stale.unlink()

    written_files: list[Path] = []
    by_subject: dict[str, list[Entry]] = {}
    for entry in entries:
        by_subject.setdefault(entry.适用学科, []).append(entry)

    for subject, group in by_subject.items():
        slug = SUBJECT_SLUGS[subject]
        path = ENTRIES_DIR / f"{slug}.yaml"
        path.write_text(dump_entries_yaml(group, subject, slug), encoding="utf-8", newline="\n")
        written_files.append(path)

    # --- 写 index/rules.json ---
    INDEX_DIR.mkdir(parents=True, exist_ok=True)
    rules_path = INDEX_DIR / "rules.json"
    write_json(rules_path, rules, compact_number_arrays=True)
    try:
        verify_rules_json(rules_path, len(entries), EMBEDDING_DIM)
    except (ValueError, json.JSONDecodeError) as exc:
        # 产物不合格就不要留在磁盘上冒充可用产物 —— 半份坏索引比没有索引更难查
        rules_path.unlink(missing_ok=True)
        print(f"写出的 rules.json 未通过自检，已删除：{exc}", file=sys.stderr)
        return 1
    written_files.append(rules_path)

    # --- 报告 ---
    scores = [score_entry(entry) for entry in entries]
    similarity = analyze_similarity(
        entries, embeddings, NEAR_DUPLICATE_COSINE[args.embedder]
    )

    # 阈值够不着时把这件事本身报成警告。一个常年沉默的检查看起来像「没问题」，
    # 而那正是它最危险的地方 —— 不能让它静悄悄地失效。
    if (
        similarity["分布"]["对数"] > 0
        and not similarity["疑似近重复"]
        and similarity["分布"]["最大"] < similarity["阈值"]
    ):
        diagnostics.append(
            Diagnostic(
                "warning",
                "（全局）",
                None,
                "",
                "近重复阈值",
                f"当前后端下同科条目两两余弦最大只有 {similarity['分布']['最大']}，"
                f"低于阈值 {similarity['阈值']}，该检查无法生效。"
                f"请按报告第八节的分布重定 —— 语义模型与哈希后端的余弦尺度差很多",
            )
        )
        # 报告与摘要都读 warnings，这里重新取一次，让上面这条也进得去
        warnings = [item for item in diagnostics if item.level == "warning"]
    self_recall = (
        run_self_recall(entries, rules["entries"], args.embedder)
        if args.self_recall
        else None
    )

    taxonomy = build_alias_taxonomy(entries)
    taxonomy_path = INDEX_DIR / "field-alias-taxonomy.json"
    write_json(taxonomy_path, taxonomy)
    written_files.append(taxonomy_path)

    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    report = {
        "builtAt": built_at,
        "model": model,
        "条目总数": len(entries),
        "错误": [item.format() for item in errors],
        "警告": [item.format() for item in warnings],
        "待核对条目": [
            {"id": entry.id, "标题": entry.标题, "来源": entry.来源}
            for entry in entries
            if entry.核验状态 == "待核对"
        ],
        "入库质量分": scores,
        "相似度": similarity,
        "自召回": self_recall,
        "产物": [str(path.relative_to(ROOT)) for path in written_files],
    }
    report_json = REPORT_DIR / "build-report.json"
    write_json(report_json, report)
    written_files.append(report_json)

    report_md = REPORT_DIR / "build-report.md"
    report_md.write_text(
        render_report(entries, diagnostics, scores, similarity, taxonomy, rules, self_recall),
        encoding="utf-8",
        newline="\n",
    )
    written_files.append(report_md)

    # --- 控制台摘要 ---
    print(f"构建完成：{len(entries)} 条规则，{len(warnings)} 条警告。")
    for subject, group in sorted(by_subject.items()):
        print(f"  {subject}: {len(group)} 条 → entries/{SUBJECT_SLUGS[subject]}.yaml")
    print(f"  向量后端: {model}")
    if self_recall:
        print(
            f"  自召回: Hit@1 {self_recall['Hit@1']}　"
            f"Hit@{TOP_K} {self_recall['Hit@%d' % TOP_K]}　MRR {self_recall['MRR']}"
        )
    mentioned = [f"{item.sheet}" for item in warnings]
    print(f"  报告: reports/build-report.md（{len(mentioned)} 条警告明细在内）")
    return 0


if __name__ == "__main__":
    # Windows 控制台默认按本机代码页解码，中文会变乱码。
    # 显式改成 UTF-8，让报错信息里的中文能原样读出来。
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8", errors="replace")
        except (AttributeError, ValueError):  # pragma: no cover - 非标准流
            pass
    sys.exit(main(sys.argv[1:]))
